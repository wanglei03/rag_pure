"""文档加载:支持 md / txt / pdf / docx,pdf 与 docx 为可选依赖(未安装时跳过并提示)。"""
from __future__ import annotations

from dataclasses import dataclass, field
from pathlib import Path


@dataclass
class Document:
    doc_id: str  # 文件名去扩展名,作为全局唯一文档标识
    source: str  # 原始文件名
    text: str
    metadata: dict = field(default_factory=dict)


def load_documents(docs_dir: Path) -> list[Document]:
    docs_dir = Path(docs_dir)
    if not docs_dir.exists():
        raise FileNotFoundError(f"文档目录不存在: {docs_dir}")

    loaders = {".md": _load_text, ".txt": _load_text, ".pdf": _load_pdf, ".docx": _load_docx}
    docs: list[Document] = []
    for path in sorted(p for p in docs_dir.iterdir() if p.is_file()):
        loader = loaders.get(path.suffix.lower())
        if loader is None:
            print(f"[跳过] 不支持的格式: {path.name}")
            continue
        text = loader(path)
        if not text.strip():
            print(f"[跳过] 空文件: {path.name}")
            continue
        docs.append(Document(doc_id=path.stem, source=path.name, text=text.strip(), metadata={"suffix": path.suffix}))
    return docs


def _load_text(path: Path) -> str:
    for enc in ("utf-8", "gbk"):
        try:
            return path.read_text(encoding=enc)
        except UnicodeDecodeError:
            continue
    return path.read_text(encoding="utf-8", errors="ignore")


def _load_pdf(path: Path) -> str:
    try:
        from pypdf import PdfReader
    except ImportError:
        print("[提示] 未安装 pypdf,跳过 PDF。pip install pypdf 后重试。")
        return ""
    reader = PdfReader(str(path))
    return "\n".join(page.extract_text() or "" for page in reader.pages)


def _load_docx(path: Path) -> str:
    try:
        import docx
    except ImportError:
        print("[提示] 未安装 python-docx,跳过 DOCX。pip install python-docx 后重试。")
        return ""
    d = docx.Document(str(path))
    return "\n".join(p.text for p in d.paragraphs)
